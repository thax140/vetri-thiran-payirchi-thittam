# ai_core/generator.py - text cleaning and output formatting (.docx, .pdf, html)
import html
import io
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

from config import FOOTER_TEXT, LOGO_PATH

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2022": "-", "\u2026": "...",
    "\u00a0": " ",
}


def sanitize_text(text: str) -> str:
    """Remove typographic quotes, markdown symbols and odd characters."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)      # **bold** -> bold
    text = re.sub(r"^#+\s*", "", text, flags=re.M)     # ## Heading -> Heading
    text = text.replace("*", "")
    return text.strip()


def _is_heading(line: str) -> bool:
    line = line.strip()
    if not line or len(line) > 70:
        return False
    return bool(re.match(r"^\d+\.\s", line)) or line.endswith(":") or line.isupper()


def _extract_terms(text: str):
    """Split semicolon-separated terms for the Terms table."""
    parts = [p.strip(" -") for p in text.split(";") if p.strip()]
    return parts if len(parts) > 1 else []


# ---------------- DOCX ----------------
def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    if os.path.exists(LOGO_PATH):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(LOGO_PATH, width=Inches(2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run(doc_type.upper())
    r.bold = True
    r.font.size = Pt(16)

    for line in text.split("\n"):
        if not line.strip():
            continue
        para = doc.add_paragraph()
        run = para.add_run(line.strip())
        if _is_heading(line):
            run.bold = True
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    term_list = _extract_terms(terms)
    if term_list:
        doc.add_paragraph().add_run("Summary of Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, t in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = t

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = FOOTER_TEXT
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- PDF ----------------
def _latin(text: str) -> str:
    # built-in PDF fonts only support latin-1
    return text.encode("latin-1", "replace").decode("latin-1")


class _LegalPDF(FPDF):
    def __init__(self, doc_type):
        super().__init__()
        self.doc_type = _latin(doc_type)

    def header(self):
        if os.path.exists(LOGO_PATH):
            self.image(LOGO_PATH, x=(self.w - 35) / 2, y=8, w=35)
            self.set_y(30)
        self.set_font("Arial", "B", 13)
        self.set_x(self.l_margin)
        self.cell(0, 8, self.doc_type, align="C")
        self.ln(12)

    def footer(self):
        self.set_y(-15)
        self.set_font("Arial", "I", 8)
        self.cell(0, 10, _latin(FOOTER_TEXT), align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = _LegalPDF(doc_type)
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    for line in text.split("\n"):
        line = _latin(line.strip())
        if not line:
            pdf.ln(3)
            continue
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Arial", "B" if _is_heading(line) else "", 11)
        pdf.multi_cell(0, 7, line)

    out = pdf.output(dest="S")
    return out.encode("latin-1") if isinstance(out, str) else bytes(out)


# ---------------- HTML preview ----------------
def format_html_preview(text: str) -> str:
    blocks = []
    for line in text.split("\n"):
        if not line.strip():
            continue
        safe = html.escape(line.strip())
        if _is_heading(line):
            blocks.append(f"<h4 style='margin:14px 0 4px;color:#fff'>{safe}</h4>")
        else:
            blocks.append(f"<p style='margin:4px 0;color:#ddd'>{safe}</p>")
    return "".join(blocks)
